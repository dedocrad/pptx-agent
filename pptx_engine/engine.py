from pptx import Presentation
from pptx.util import Inches
from pptx.dml.color import RGBColor

from pptx_engine.layouts import RENDERERS, SLIDE_W, SLIDE_H
from pptx_engine.spec import PresentationSpec
from pptx_engine.theme import THEMES


def _set_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(*color)


def build_presentation(spec: PresentationSpec) -> Presentation:
    theme = THEMES.get(spec.theme, THEMES["default"])

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    blank = prs.slide_layouts[6]  # empty layout

    for slide_spec in spec.slides:
        slide = prs.slides.add_slide(blank)
        _set_bg(slide, theme.bg_color)

        renderer = RENDERERS.get(slide_spec.layout)
        if renderer is None:
            raise ValueError(f"Unknown layout: {slide_spec.layout}")
        renderer(slide, slide_spec, theme)

    return prs


def save(spec: PresentationSpec, path: str):
    prs = build_presentation(spec)
    prs.save(path)
    return path