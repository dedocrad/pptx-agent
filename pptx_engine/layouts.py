from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

from pptx_engine.spec import SlideSpec
from pptx_engine.theme import Theme

SLIDE_W = Inches(13.333)   # 16:9
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.6)
TITLE_TOP = MARGIN
TITLE_HEIGHT = Inches(0.9)
CONTENT_TOP = Inches(1.8)
CONTENT_HEIGHT = Inches(5)
COLUMN_WIDTH = (SLIDE_W - 3 * MARGIN) / 2
COLUMN_RIGHT = 2 * MARGIN + COLUMN_WIDTH


def _rgb(t): return RGBColor(*t)


def _add_textbox(slide, left, top, width, height,
                 text, theme: Theme, size=None, bold=False,
                 color=None, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.text = text
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.runs[0]
    run.font.name = theme.body_font
    run.font.size = Pt(size or theme.body_size)
    run.font.bold = bold
    run.font.color.rgb = _rgb(color or theme.text_color)
    return box


def _add_standard_title(slide, title: str, theme: Theme):
    _add_textbox(
        slide,
        MARGIN,
        TITLE_TOP,
        SLIDE_W - 2 * MARGIN,
        TITLE_HEIGHT,
        title,
        theme,
        size=theme.title_size,
        bold=True,
        color=theme.title_color,
    )


def _add_bullets(slide, left, top, width, height,
                 items, theme: Theme, size=None):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"• {item}"
        p.font.name = theme.body_font
        p.font.size = Pt(size or theme.body_size)
        p.font.color.rgb = _rgb(theme.text_color)
    return box


# ---------- layouts ----------

def render_title(slide, s: SlideSpec, theme: Theme):
    _add_textbox(slide, MARGIN, Inches(2.5), SLIDE_W - 2*MARGIN, Inches(1.2),
                 s.title, theme, size=44, bold=True,
                 color=theme.title_color, align=PP_ALIGN.CENTER)
    if s.subtitle:
        _add_textbox(slide, MARGIN, Inches(3.8), SLIDE_W - 2*MARGIN, Inches(0.8),
                     s.subtitle, theme, size=20,
                     color=theme.text_color, align=PP_ALIGN.CENTER)


def render_text(slide, s: SlideSpec, theme: Theme):
    _add_standard_title(slide, s.title, theme)
    _add_bullets(
        slide, MARGIN, CONTENT_TOP, SLIDE_W - 2 * MARGIN, CONTENT_HEIGHT,
        s.bullets, theme,
    )


def render_image(slide, s: SlideSpec, theme: Theme):
    _add_standard_title(slide, s.title, theme)
    if s.image:
        slide.shapes.add_picture(
            s.image, MARGIN, CONTENT_TOP, width=SLIDE_W - 2 * MARGIN,
        )


def render_text_image(slide, s: SlideSpec, theme: Theme):
    _add_standard_title(slide, s.title, theme)
    _add_bullets(
        slide, MARGIN, CONTENT_TOP, COLUMN_WIDTH, CONTENT_HEIGHT, s.bullets,
        theme,
    )
    if s.image:
        slide.shapes.add_picture(
            s.image, COLUMN_RIGHT, CONTENT_TOP, width=COLUMN_WIDTH,
        )


def render_image_text(slide, s: SlideSpec, theme: Theme):
    _add_standard_title(slide, s.title, theme)
    if s.image:
        slide.shapes.add_picture(
            s.image, MARGIN, CONTENT_TOP, width=COLUMN_WIDTH,
        )
    _add_bullets(
        slide, COLUMN_RIGHT, CONTENT_TOP, COLUMN_WIDTH, CONTENT_HEIGHT,
        s.bullets, theme,
    )


def render_two_columns(slide, s: SlideSpec, theme: Theme):
    _add_standard_title(slide, s.title, theme)
    _add_bullets(
        slide, MARGIN, CONTENT_TOP, COLUMN_WIDTH, CONTENT_HEIGHT,
        s.column_left, theme,
    )
    _add_bullets(
        slide, COLUMN_RIGHT, CONTENT_TOP, COLUMN_WIDTH, CONTENT_HEIGHT,
        s.column_right, theme,
    )


def render_references(slide, s: SlideSpec, theme: Theme):
    _add_standard_title(slide, s.title or "Источники", theme)
    _add_bullets(
        slide, MARGIN, CONTENT_TOP, SLIDE_W - 2 * MARGIN, CONTENT_HEIGHT,
        s.references, theme, size=14,
    )


RENDERERS = {
    "title": render_title,
    "text": render_text,
    "image": render_image,
    "text_image": render_text_image,
    "image_text": render_image_text,
    "two_columns": render_two_columns,
    "references": render_references,
}